import streamlit as st
import pandas as pd
import io
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class BaseStreamlitView:

    @staticmethod
    def show_error(message: str):
        # st.error(message)
        st.toast(message, icon="🚨")

    @staticmethod
    def show_warning(message: str):
        # st.warning(message)
        st.toast(message, icon="⚠️")

    @staticmethod
    def show_success(message: str):
        st.toast(message, icon="✔️")

    def download_data_display(self, filtered_df: pd.DataFrame, key="data_download"):

        # Optional: Download filtered data
        if not filtered_df.empty:
            # csv = filtered_df.to_csv(index=False).encode('utf-8')
            # 1. Chuyển DataFrame thành chuỗi CSV (chưa encode)
            csv_string = filtered_df.to_csv(index=False)
             # 2. Thêm ký tự BOM '\ufeff' vào đầu chuỗi rồi mới encode sang utf-8
            csv_bytes = ("\ufeff" + csv_string).encode("utf-8")
            st.download_button(
                label="Download Displayed Data as CSV",
                data=csv_bytes,
                key=key,
                file_name=f"{key}_{pd.Timestamp.now().strftime('%Y%m%d%H%M')}.csv",
                mime="text/csv",
                icon=":material/download:",
            )

        
    def download_data_excel_display(self, df, key="data_download"):
         if not df.empty:
            # 1. Tạo một buffer ảo trong bộ nhớ để ghi file Excel
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Report')
            
            # 2. Lấy dữ liệu dạng bytes từ buffer
            excel_bytes = buffer.getvalue()
            
            # 3. Tạo nút download với định dạng excel
            st.download_button(
                label="Download Displayed Data as Excel",
                data=excel_bytes,
                file_name=f"{key}_{pd.Timestamp.now().strftime('%Y%m%d%H%M')}.xlsx", # Đổi đuôi file thành .xlsx
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=key
            )



    def download_data_excel_format_display(self, df, key="data_download"):
        if not df.empty:
            # 1. Tạo một file Excel ảo trong bộ nhớ
            buffer = io.BytesIO()
            
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                # XUẤT DỮ LIỆU ĐÈ LÊN SÁT DÒNG A2 (startrow=1 trong pandas tương đương dòng 2 Excel)
                df.to_excel(writer, index=False, sheet_name='Report', startrow=1)
                
                workbook = writer.book
                worksheet = writer.sheets['Report']
                
                # ÉP EXCEL BẬT CHẾ ĐỘ XEM TRANG IN (Để người dùng thấy ngay tiêu đề lặp lại)
                # worksheet.sheet_view.view = "pageLayout"
                
                num_cols = len(df.columns)
                last_col_letter = get_column_letter(num_cols)
                
                # -------------------------------------------------------------
                # BƯỚC 2: ĐỔI THÀNH KHỔ GIẤY DỌC A4 TRÀN LỀ
                # -------------------------------------------------------------
                worksheet.page_setup.paperSize = 9  # Khổ A4
                worksheet.page_setup.orientation = 'portrait'  # ĐỊNH DẠNG KHỔ GIẤY DỌC
                # worksheet.page_setup.orientation = 'landscape' # Khổ ngang

                # Thiết lập tự động căn giữa trang theo chiều ngang (Center horizontally)
                worksheet.print_options.horizontalCentered = True
                
                # Thu lề tối đa (Narrow Margins)
                worksheet.page_margins.left = 0.25    
                worksheet.page_margins.right = 0.25   
                worksheet.page_margins.top = 0.5      
                worksheet.page_margins.bottom = 0.5   
                worksheet.page_margins.header = 0.2
                worksheet.page_margins.footer = 0.2

                # Số trang động "Trang X / Y" ở chính giữa chân trang
                worksheet.oddFooter.center.text = "Trang &P / &N"
                worksheet.oddFooter.center.size = 10
                worksheet.oddFooter.center.font = "Calibri"

                # Ép vừa khít chiều ngang khổ dọc (không bị vỡ cột sang trang khác)
                worksheet.sheet_properties.pageSetUpPr.fitToPage = True
                worksheet.page_setup.fitToWidth = 1   
                worksheet.page_setup.fitToHeight = 0  

                # -------------------------------------------------------------
                # BƯỚC 3: SỬA LỖI LẶP TIÊU ĐỀ - KHÓA DÒNG 1 VÀ 2 TRÊN TẤT CẢ CÁC TRANG
                # -------------------------------------------------------------
                # Thuộc tính này nạp trực tiếp Print Titles chuẩn Excel. 
                # Khóa toàn bộ vùng từ dòng 1 đến dòng 2 (gồm cả Tiêu đề lớn & Tên cột)
                worksheet.print_title_rows = "1:2"

                # -------------------------------------------------------------
                # BƯỚC 4: GỘP Ô DÒNG TIÊU ĐỀ 1 VÀ CĂN GIỮA TOÀN BẢNG
                # -------------------------------------------------------------
                # "BÁO CÁO CHI TIẾT BIN TO BIN MOVEMENT"
                title_text = "LIST FG LEVEL A NEED BIN TO BIN"
                worksheet['A1'] = title_text
                
                # Gộp từ cột A sang cột cuối cùng trên dòng 1
                worksheet.merge_cells(f"A1:{last_col_letter}1")
                
                title_font = Font(name="Calibri", size=22, bold=True, color="FF000000") #1F4E78
                worksheet['A1'].font = title_font
                worksheet['A1'].alignment = Alignment(horizontal="center", vertical="center")
                worksheet.row_dimensions[1].height = 30  # Chiều cao rộng rãi cho tiêu đề gộp dòng 1
                worksheet.row_dimensions[2].height = 24  # Chiều cao cho dòng tên cột dữ liệu dòng 2
                
                # -------------------------------------------------------------
                # BƯỚC 5: ĐÒNG KHUNG (BORDER) TOÀN BỘ CỘT VÀ DÒNG DATA
                # -------------------------------------------------------------
                thin_side = Side(border_style="thin", color="A6A6A6")
                data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
                
                header_font = Font(name="Calibri", size=11, bold=True, color="000000")
                data_font = Font(name="Calibri", size=13)
                
                # Lặp quét từ dòng số 2 (Tên cột dữ liệu) đổ xuống
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                    # Lấy số thứ tự dòng hiện tại
                    current_row_idx = row[0].row
                    # Nếu là dòng dữ liệu (từ dòng 3 trở đi), đặt chiều cao dòng là 20
                    if current_row_idx >= 3:
                        worksheet.row_dimensions[current_row_idx].height = 20

                    for cell in row:
                        cell.border = data_border
                        if cell.row == 2:
                            cell.font = header_font
                            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        else:
                            cell.font = data_font
                            if isinstance(cell.value, (int, float)):
                                cell.alignment = Alignment(horizontal="right", vertical="center")
                            else:
                                cell.alignment = Alignment(horizontal="left", vertical="center")

                # -------------------------------------------------------------
                # BƯỚC 6: TỰ ĐỘNG GIÃN RỘNG CỘT THEO CHIỀU DÀI CHỮ
                # -------------------------------------------------------------
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    
                    for cell in col:
                        if cell.row == 1: 
                            continue  # Bỏ qua dòng 1 gộp ô để không tính sai chiều rộng cột
                        if cell.value is not None:
                            max_len = max(max_len, len(str(cell.value)))
                    
                    # Khổ dọc chừa hao lề khít (+4) để giữ độ co giãn tối đa cho trang giấy
                    worksheet.column_dimensions[col_letter].width = max(max_len + 4, 9)

            # 7. Lấy dữ liệu bytes hoàn chỉnh từ file Excel ảo
            excel_bytes = buffer.getvalue()
            
            # 8. Tạo nút bấm download trên Streamlit giao diện Excel (.xlsx)
            st.download_button(
                label="Download Displayed Data as Excel",
                data=excel_bytes,
                file_name=f"{key}_{pd.Timestamp.now().strftime('%Y%m%d%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=key
            )


    def search_df(self, filtered_df: pd.DataFrame, search_query: str) -> pd.DataFrame:
        if search_query:
            mask = pd.Series(False, filtered_df.index)
            for col in filtered_df.columns:
                mask |= filtered_df[col].astype(str).str.contains(search_query.strip(), case=False, na=False)
            filtered_df = filtered_df[mask]
        return filtered_df



#=================================================================================================
# import io
# import pandas as pd
# import streamlit as st
# from openpyxl.styles import Font, Alignment, Border, Side
# from openpyxl.utils import get_column_letter
# # Download data tự động thêm dòng thiêu đề và đóng khung
# def download_data_display(self, df, key="download_btn"):
#     # 1. Tạo một file Excel ảo trong bộ nhớ
#     buffer = io.BytesIO()
    
#     with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
#         # Xuất dữ liệu từ dòng số 3 (startrow=2 trong Python vì đếm từ 0) để chừa khoảng trống cho tiêu đề lớn
#         df.to_excel(writer, index=False, sheet_name='Report', startrow=2)
        
#         # Lấy workbook và sheet hiện tại để định dạng
#         workbook = writer.book
#         worksheet = writer.sheets['Report']
        
#         # -------------------------------------------------------------
#         # BƯỚC 2: THÊM DÒNG TIÊU ĐỀ LỚN CHO BẢO CÁO
#         # -------------------------------------------------------------
#         title_text = "BÁO CÁO CHI TIẾT BIN TO BIN MOVEMENT" # Bạn có thể đổi tên tiêu đề ở đây
#         worksheet['A1'] = title_text
        
#         # Định dạng cho tiêu đề (Chữ đậm, size 16, màu xanh đậm)
#         title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
#         worksheet['A1'].font = title_font
        
#         # -------------------------------------------------------------
#         # BƯỚC 3: ĐỊNH DẠNG ĐÓNG KHUNG (BORDER) VÀ CHỮ CHO TOÀN BẢNG
#         # -------------------------------------------------------------
#         # Định nghĩa nét viền mỏng (Thin border)
#         thin_side = Side(border_style="thin", color="D9D9D9") # Màu xám nhạt cho tinh tế
#         data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
#         # Kiểu chữ tiêu đề cột (Dòng số 3)
#         header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
#         # Bạn có thể đổ màu nền cho Header nếu muốn (Tùy chọn)
#         # from openpyxl.styles import PatternFill
#         # header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
#         # Kiểu chữ cho phần nội dung data
#         data_font = Font(name="Calibri", size=11)
        
#         # Lặp qua tất cả các ô có chứa dữ liệu để đóng khung và căn chỉnh
#         for row in worksheet.iter_rows(min_row=3, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
#             for cell in row:
#                 cell.border = data_border  # Đóng khung cho ô này
                
#                 if cell.row == 3:
#                     # Nếu là dòng tiêu đề cột
#                     cell.font = header_font
#                     cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#                     # cell.fill = header_fill # Bỏ dấu thăng nếu muốn đổ màu nền header
#                 else:
#                     # Nếu là dòng dữ liệu bình thường
#                     cell.font = data_font
#                     # Nếu là cột chứa chữ thì căn trái, cột chứa số thì căn phải
#                     if isinstance(cell.value, (int, float)):
#                         cell.alignment = Alignment(horizontal="right", vertical="center")
#                     else:
#                         cell.alignment = Alignment(horizontal="left", vertical="center")

#         # -------------------------------------------------------------
#         # BƯỚC 4: TỰ ĐỘNG CĂN RỘNG CỘT (Để không bị lỗi hiển thị ###)
#         # -------------------------------------------------------------
#         for col in worksheet.columns:
#             max_len = 0
#             col_letter = get_column_letter(col[0].column)
#             for cell in col:
#                 # Bỏ qua dòng tiêu đề lớn dòng 1 khi tính độ rộng cột
#                 if cell.row == 1: 
#                     continue
#                 if cell.value:
#                     max_len = max(max_len, len(str(cell.value)))
#             # Đặt độ rộng cột bằng chiều dài text + 4 khoảng trắng trừ hao
#             worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

#     # 5. Lấy dữ liệu bytes hoàn chỉnh từ file Excel ảo
#     excel_bytes = buffer.getvalue()
    
#     # 6. Tạo nút bấm download trên Streamlit giao diện Excel (.xlsx)
#     st.download_button(
#         label="📥 Tải Báo Cáo Gồm Tiêu Đề & Khung Bảng",
#         data=excel_bytes,
#         file_name="bin_to_bin_report.xlsx", # Đổi tên file đích tại đây
#         mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#         key=key
#     )
#===============================================================================================
